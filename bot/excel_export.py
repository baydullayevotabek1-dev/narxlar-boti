"""Build the Excel workbook managers download after a search."""
import datetime
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import EZVIZ_SECONDARY_DISCOUNT

HEAD_FILL = PatternFill("solid", fgColor="667EEA")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="333333")
SUB_FONT = Font(size=10, color="888888")
BEST_FILL = PatternFill("solid", fgColor="D5F5E3")
BEST_FONT = Font(bold=True, color="1E7E34")
THIN = Side(style="thin", color="DDDDDD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = "#,##0.00"


def _rows_for(store: str, r: dict) -> list[dict]:
    """One spreadsheet row per priced option. Ezviz quotes two discounts."""
    base = {
        "store": store,
        "price": r["price"],
        "manual": bool(r.get("is_override")),
    }
    if store.lower() == "ezviz":
        return [
            {**base, "discount": r["discount"], "final": round(r["price"] * (1 - r["discount"] / 100), 2)},
            {**base, "discount": EZVIZ_SECONDARY_DISCOUNT,
             "final": round(r["price"] * (1 - EZVIZ_SECONDARY_DISCOUNT / 100), 2)},
        ]
    return [{**base, "discount": r["discount"], "final": r["final"]}]


def _autosize(ws, widths: list[int]):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_header(ws, headers: list[str], row: int):
    for col, name in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=name)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER


def _sheet_detail(ws, queries, found, store_dates, stamp):
    ws.cell(row=1, column=1, value="Narxlar solishtiruvi — batafsil").font = TITLE_FONT
    ws.cell(row=2, column=1, value=f"Eksport: {stamp}").font = SUB_FONT

    headers = ["Qidirilgan", "Model", "Turi", "Do'kon",
               "Asl narx ($)", "Skidka (%)", "Yakuniy narx ($)", "Narx manbasi",
               "Narx yangilangan"]
    _write_header(ws, headers, 4)
    _autosize(ws, [24, 28, 14, 16, 13, 11, 15, 16, 18])

    r = 5
    for q in queries:
        fam = found.get(q)
        if not fam:
            continue
        for g in fam["groups"]:
            kind = "asosiy" if g["is_base"] else "qo'shimchali"
            if g["is_asked"] and not g["is_base"]:
                kind = "so'ralgan"
            best = min((x["final"] for x in g["rows"]), default=None)
            for res in g["rows"]:
                for row in _rows_for(res["store"], res):
                    is_best = len(g["rows"]) > 1 and best is not None and abs(row["final"] - best) < 0.01
                    values = [
                        q, g["model"], kind, row["store"],
                        row["price"], row["discount"] / 100, row["final"],
                        "qo'lda tuzatilgan" if row["manual"] else "fayldan",
                        store_dates.get(row["store"], ""),
                    ]
                    for col, v in enumerate(values, start=1):
                        c = ws.cell(row=r, column=col, value=v)
                        c.border = BORDER
                        if col in (5, 7):
                            c.number_format = MONEY
                        if col == 6:
                            c.number_format = "-0%"
                        if is_best:
                            c.fill = BEST_FILL
                    if g["is_base"]:
                        ws.cell(row=r, column=2).font = Font(bold=True)
                    if is_best:
                        ws.cell(row=r, column=7).font = BEST_FONT
                    r += 1
        r += 1  # blank line between searches

    ws.freeze_panes = "A5"


def _sheet_compare(ws, queries, found, stamp):
    ws.cell(row=1, column=1, value="Narxlar solishtiruvi — jadval").font = TITLE_FONT
    ws.cell(row=2, column=1, value=f"Eksport: {stamp}").font = SUB_FONT

    # One row per distinct model across all searches, in search order.
    models: list[dict] = []
    seen = set()
    for q in queries:
        fam = found.get(q)
        if not fam:
            continue
        for g in fam["groups"]:
            if g["model"] not in seen:
                seen.add(g["model"])
                models.append(g)

    stores = sorted({r["store"] for g in models for r in g["rows"]})
    headers = ["Model"] + stores + ["Eng arzon", "Eng arzon narx ($)"]
    _write_header(ws, headers, 4)
    _autosize(ws, [30] + [14] * len(stores) + [18, 18])

    r = 5
    for g in models:
        mc = ws.cell(row=r, column=1, value=g["model"])
        mc.border = BORDER
        if g["is_base"]:
            mc.font = Font(bold=True)

        by_store = {x["store"]: x for x in g["rows"]}
        for col, store in enumerate(stores, start=2):
            c = ws.cell(row=r, column=col)
            c.border = BORDER
            hit = by_store.get(store)
            if not hit:
                c.value = "—"
                c.alignment = Alignment(horizontal="center")
                continue
            c.value = hit["final"]
            c.number_format = MONEY

        winner = min(g["rows"], key=lambda x: x["final"])
        wc = ws.cell(row=r, column=len(stores) + 2, value=winner["store"])
        wp = ws.cell(row=r, column=len(stores) + 3, value=winner["final"])
        wp.number_format = MONEY
        for c in (wc, wp):
            c.border = BORDER
            c.fill = BEST_FILL
            c.font = BEST_FONT
        r += 1

    ws.freeze_panes = "B5"

    note = ws.cell(row=r + 1, column=1,
                   value="Qalin — asosiy model. Qolganlari shu modelning qo'shimchali variantlari.")
    note.font = SUB_FONT


def _sheet_missing(ws, suggestions, not_found):
    ws.cell(row=1, column=1, value="Topilmagan modellar").font = TITLE_FONT
    _write_header(ws, ["Qidirilgan model", "Holat", "Taklif qilingan modellar"], 3)
    _autosize(ws, [26, 20, 70])

    r = 4
    for q, names in suggestions.items():
        ws.cell(row=r, column=1, value=q).border = BORDER
        ws.cell(row=r, column=2, value="aniq topilmadi").border = BORDER
        ws.cell(row=r, column=3, value=", ".join(names)).border = BORDER
        r += 1
    for q in not_found:
        ws.cell(row=r, column=1, value=q).border = BORDER
        ws.cell(row=r, column=2, value="umuman topilmadi").border = BORDER
        ws.cell(row=r, column=3, value="—").border = BORDER
        r += 1


def build_workbook(queries, found, suggestions, not_found, store_dates) -> bytes:
    stamp = datetime.datetime.now().strftime("%d.%m.%Y %H:%M")
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Batafsil"
    _sheet_detail(ws1, queries, found, store_dates, stamp)

    ws2 = wb.create_sheet("Solishtirish")
    _sheet_compare(ws2, queries, found, stamp)

    only_missing = {q: n for q, n in suggestions.items() if q not in found}
    if only_missing or not_found:
        _sheet_missing(wb.create_sheet("Topilmagan"), only_missing, not_found)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
